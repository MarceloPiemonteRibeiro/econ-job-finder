#!/usr/bin/env python3
"""Build master_sources.xlsx = sources.yaml + career-center Google Sheet + IHEID Resources PDF,
deduped by URL, classified by category. One-off skeleton builder; re-run to refresh."""
import csv, re
from pathlib import Path
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = Path(__file__).resolve().parent
S = BASE / "data"  # put sheet.csv and resources.pdf here before running
OUT = BASE / "master_sources.xlsx"
AUTO = {"imf", "banque de france", "oecd", "ebrd", "ceb", "aiib", "unjobs", "jobs.ac.uk", "inomics", "euraxess", "impactpool"}


def norm(u):
    u = re.sub(r"^https?://", "", (u or "").strip(), flags=re.I)
    u = re.sub(r"^www\.", "", u, flags=re.I)
    return u.rstrip("/").lower()


def unify(x):
    x = (x or "").lower()
    checks = [
        ("Central banks", ["central bank"]),
        ("Job boards & aggregators", ["job board", "job search", "aggregator", "academic aggregator"]),
        ("Consulting", ["consulting", "consultanc", "management consult"]),
        ("Sustainable finance & impact", ["sustainable finance", "impact invest"]),
        ("Private sector / finance", ["private sector", "banking", "asset manage", "insurance", "commodit", "investment bank", "rating"]),
        ("Policy / IO / IFIs", ["central bank" if False else "multilateral", "ifi", "intergovernmental", "international org", "policy / io", "policy/io", "un system", "development bank", "un specialized", "un secretariat", "un regional"]),
        ("Development & cooperation", ["cooperation and development", "development", "humanitarian"]),
        ("Environment & sustainability", ["environment", "sustainab", "climate"]),
        ("Think tanks & research / Academic", ["think tank", "think thank", "research and think", "research center", "research institute", "academic", "university"]),
        ("Social enterprise & innovation", ["social entrepreneur", "innovation"]),
        ("Swiss public sector", ["swiss public", "federal administration", "cantonal", "public sector", "government"]),
        ("Other sectors (advocacy/health/rights/law/media)", ["advocacy", "conflict", "gender", "health", "human and children", "rights", "international law", "media", "law"]),
    ]
    for cat, keys in checks:
        if any(k in x for k in keys):
            return cat
    return "Other"


def parse_sources():
    out, cat, sub = [], "Other", ""
    for ln in (BASE / "sources.yaml").read_text(encoding="utf-8").splitlines():
        s = ln.strip()
        if s.startswith("#"):
            txt = s.lstrip("#").strip()
            if txt and "═" not in txt and "─" not in txt and 2 < len(txt) < 80 and not re.match(
                    r"(type|url|selector|credentials|note|for |annotation|rss|html|manual|\[|http|ponytail|--)", txt, re.I):
                cat, sub = ([x.strip() for x in txt.split("—", 1)] if "—" in txt else [txt, ""])
            continue
        m = re.search(r'name:\s*"([^"]+)".*?url:\s*"([^"]+)"(?:.*?note:\s*"([^"]*)")?', s)
        if m:
            out.append({"category": unify(cat), "subcategory": sub, "org": m.group(1), "url": m.group(2),
                        "location": "", "type": "", "roles": "", "notes": m.group(3) or "", "origins": {"watcher-list"}})
    return out


def parse_sheet():
    out = []
    with open(S / "sheet.csv", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            org = (row.get("Organization") or "").strip()
            if not org:
                continue
            typ = (row.get("Type") or "").strip()
            roles = (row.get("Typical PhD\u2011Econ Roles") or row.get("Typical PhD-Econ Roles") or "").strip()
            out.append({"category": unify(typ), "subcategory": "", "org": org, "url": (row.get("Careers URL") or "").strip(),
                        "location": (row.get("City") or "").strip(), "type": typ, "roles": roles,
                        "notes": " | ".join(x for x in [(row.get("Hiring Departments") or "").strip(), (row.get("Comments") or "").strip()] if x),
                        "origins": {"career-center-sheet"}})
    return out


SECTIONS = ["Job Search Sites", "Advocacy and General Interest", "Conflict and Post Conflict Management", "Consulting",
            "Environment and Sustainability", "Gender and Diversity", "Health and Sanitation", "Human and Children Rights",
            "International Cooperation and Development", "International Law", "Media and Communication",
            "Research and Think Tanks", "Social Entrepreneurship and Innovation", "Sustainable Finance and Impact Investment",
            "Swiss Public Sector", "Federal Administration", "Cantonal Administration"]


def parse_pdf():
    txt = "\n".join((p.extract_text() or "") for p in PdfReader(str(S / "resources.pdf")).pages)
    t = re.sub(r"(https?://|www\.)\s+", r"\1", txt)   # rejoin wrapped scheme
    t = re.sub(r"[ \t]+", " ", t)
    pos = sorted((m.start(), sec) for sec in SECTIONS for m in re.finditer(re.escape(sec), t))
    out = []
    for i, (p, sec) in enumerate(pos):
        chunk = t[p:(pos[i + 1][0] if i + 1 < len(pos) else len(t))]
        for m in re.finditer(r"([A-Z][A-Za-z0-9 &\u2019'()/.,\-]{2,70}?):\s*(https?://[^\s]+|www\.[^\s]+)", chunk):
            name = re.sub(r"\s+", " ", m.group(1)).strip(" \u2022-\u25cf")
            url = m.group(2).strip().rstrip(".,);")
            url = "https://" + url if url.startswith("www.") else url
            out.append({"category": unify(sec), "subcategory": sec, "org": name[:80], "url": url,
                        "location": "", "type": "", "roles": "", "notes": "", "origins": {"IHEID-PDF"}})
    return out


# ---- merge (dedupe by URL) ----
src, sheet, pdf = parse_sources(), parse_sheet(), parse_pdf()
master = {}
def add(e):
    k = norm(e["url"]) or e["org"].lower()
    if k in master:
        m = master[k]
        m["origins"] |= e["origins"]
        for fld in ("location", "type", "roles", "notes", "subcategory"):
            if not m[fld] and e[fld]:
                m[fld] = e[fld]
        if m["category"] in ("Other", "") and e["category"] not in ("Other", ""):
            m["category"] = e["category"]
    else:
        master[k] = dict(e)

for e in src:   add(e)     # sources.yaml first (authoritative category)
for e in sheet: add(e)
for e in pdf:   add(e)

# coverage: sheet orgs present in sources.yaml (by URL)
src_keys = {norm(e["url"]) for e in src}
sheet_in = sum(1 for e in sheet if norm(e["url"]) in src_keys)
print(f"sources.yaml: {len(src)} | sheet: {len(sheet)} ({sheet_in} already in sources.yaml, {len(sheet)-sheet_in} new) | PDF: {len(pdf)} | MASTER (deduped): {len(master)}")

# ---- write xlsx ----
CAT_ORDER = ["Job boards & aggregators", "Central banks", "Policy / IO / IFIs", "Development & cooperation",
             "Think tanks & research / Academic", "Consulting", "Private sector / finance",
             "Sustainable finance & impact", "Environment & sustainability", "Social enterprise & innovation",
             "Swiss public sector", "Other sectors (advocacy/health/rights/law/media)", "Other"]
rows = list(master.values())
rows.sort(key=lambda r: (CAT_ORDER.index(r["category"]) if r["category"] in CAT_ORDER else 99, r["subcategory"], r["org"].lower()))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "All sources"
headers = ["Category", "Sub-category", "Organization", "URL", "Location", "Type", "Typical PhD-econ roles", "Auto-pulled?", "Notes / comments", "Origin", "Works? (verify)"]
ws.append(headers)
hdr_fill = PatternFill("solid", fgColor="1F4E78"); thin = Side(style="thin", color="DDDDDD")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=11); c.fill = hdr_fill; c.alignment = Alignment(vertical="center")
for r in rows:
    auto = "YES" if (r["org"].lower() in AUTO or any(a in r["org"].lower() for a in AUTO)) else ""
    ws.append([r["category"], r["subcategory"], r["org"], r["url"], r["location"], r["type"], r["roles"],
               auto, r["notes"], ", ".join(sorted(r["origins"])), ""])
# style
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.border = Border(bottom=thin)
    if row[7].value == "YES":
        row[7].fill = PatternFill("solid", fgColor="C6EFCE"); row[7].font = Font(bold=True, color="1a7f37")
    if row[3].value:
        row[3].hyperlink = row[3].value; row[3].font = Font(color="1155CC", underline="single")
widths = [30, 26, 34, 46, 16, 22, 30, 12, 40, 20, 12]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(OUT)
print("wrote", OUT, "rows:", len(rows))
